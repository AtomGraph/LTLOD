"""Pre-normalize the facility registers (etl/facilities/) into flat CSVs for CSV2RDF.

Three public registers, each keyed by its own natural key:

  education  ŠVIS education & science institutions (get.data.gov.lt
             lsd/svietimo_istaigos/SvietimoIstaiga, all ~10k incl. closed ones),
             enriched from nsa/svis/bendrojo_ugdymo_mokyklos/BugdymoMokykla
             (general-education schools: ownership, parent institution of branches).
  health     VASPVT licensed health-care institutions (vaspvt/licencijos/Istaiga)
             and their activity addresses (IstaigosVeiklosAdresas).
  heritage   KPD Cultural Property Register (KVR) immovable objects — an XLSX
             distribution on data.gov.lt (dataset 2192), not in Spinta.

What this step does that a SPARQL mapping cannot do sensibly:
  * reprojects ŠVIS LKS-94 points to WGS84 (⚠ EPSG:3346 WKT is "POINT (N E)" —
    northing first, see geo.py; every point is bounds-checked against Lithuania);
  * resolves free-text municipality / eldership names (KVR, VASPVT addresses) to AR
    codes against the AR Savivaldybe / Seniunija lists — exact name match only,
    unresolved names are left blank (never guessed) and counted;
  * mints ASCII slugs for derived taxonomy concepts;
  * reads the KVR XLSX.

Every output column is a plain value; the RDF is produced by etl/facilities/mappings/*.rq.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime

from pyproj import Transformer

from .geo import LAT_MAX, LAT_MIN, LON_MAX, LON_MIN


def slugify(text: str) -> str:
    """ASCII slug: strip diacritics, lowercase, non-alphanumerics -> '-'."""
    ascii_ = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_.lower()).strip("-")


def read_csv(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write_csv(path: str, fields: list[str], rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {len(rows)} rows -> {path}", file=sys.stderr)


def is_jar(code: str) -> bool:
    """A JAR (Register of Legal Entities) company code is 9 digits. ŠVIS also issues
    short internal codes to branches (skyriai) that are not legal entities."""
    return bool(re.fullmatch(r"[0-9]{9}", code or ""))


def iso_date(value) -> str:
    if value in (None, "", "NULL"):
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value)[:10]


# --- AR name lookup ------------------------------------------------------------

class AdminNames:
    """AR municipality / eldership names -> codes. AR names are bare stems
    ("Trakų rajono", "Vilniaus miesto", "Birštono"); sources append " sav." / " sen."
    and sometimes abbreviate "rajono" -> "r." and "miesto" -> "m."."""

    def __init__(self, municipalities: str, elderships: str):
        self.sav = {}
        for r in read_csv(municipalities):
            self.sav[self._norm(r["pavadinimas"])] = r["sav_kodas"]
        # Current elderships win; an ended one (sen_iki, abolished by a reform after the
        # source snapshot was taken) is the fallback, but only when its name is unique
        # within the municipality — a reused name is ambiguous and stays unresolved.
        self.sen, ended = {}, {}
        for r in read_csv(elderships):
            key = (r["savivaldybe.sav_kodas"], self._norm(r["pavadinimas"]))
            if r.get("sen_iki"):
                ended.setdefault(key, []).append(r["sen_kodas"])
            else:
                self.sen[key] = r["sen_kodas"]
        for key, codes in ended.items():
            if key not in self.sen and len(codes) == 1:
                self.sen[key] = codes[0]
        self.misses: Counter = Counter()

    @staticmethod
    def _norm(name: str) -> str:
        name = name.strip()
        name = re.sub(r"\s*\b(sav|sen)\.?$", "", name)
        name = re.sub(r"\br\.$", "rajono", name)
        name = re.sub(r"\bm\.$", "miesto", name)
        name = re.sub(r"\bmiesto\s+savivaldybė$", "miesto", name)
        return re.sub(r"\s+", " ", name)

    def municipality(self, name: str) -> str:
        if not name or name == "NULL":
            return ""
        norm = self._norm(name)
        code = self.sav.get(norm, "")
        if not code and norm.endswith(" miesto"):
            # "Marijampolės m. sav." is AR "Marijampolės" (no "miesto" in its name)
            code = self.sav.get(norm[: -len(" miesto")], "")
        if not code:
            self.misses["sav:" + name] += 1
        return code

    def eldership(self, sav_code: str, name: str) -> str:
        if not sav_code or not name or name == "NULL":
            return ""
        norm = self._norm(name)
        code = self.sen.get((sav_code, norm), "")
        if not code and norm.endswith(" miesto"):
            # a town's own eldership: KVR "Jonavos miesto sen." is AR "Jonavos"
            code = self.sen.get((sav_code, norm[: -len(" miesto")]), "")
        if not code:
            self.misses["sen:" + name] += 1
        return code

    def report(self, label: str) -> None:
        total = sum(self.misses.values())
        if total:
            print(f"{label}: {total} unresolved admin-unit names "
                  f"({len(self.misses)} distinct), e.g. {self.misses.most_common(5)}",
                  file=sys.stderr)


# --- education -----------------------------------------------------------------

POINT_RE = re.compile(r"POINT\s*\(\s*([0-9.]+)\s+([0-9.]+)\s*\)")


def education(args) -> None:
    rows = read_csv(args.institutions)
    schools = {r["ins_id"]: r for r in read_csv(args.schools)}
    to_wgs84 = Transformer.from_crs("EPSG:3346", "EPSG:4326", always_xy=True)
    out, out_of_bounds = [], []
    for r in rows:
        school = schools.get(r["ins_id"], {})
        lat = lon = ""
        m = POINT_RE.match(r.get("koord") or "")
        if m:
            northing, easting = float(m.group(1)), float(m.group(2))  # N first (LKS-94)
            lon_f, lat_f = to_wgs84.transform(easting, northing)
            if not (LAT_MIN <= lat_f <= LAT_MAX and LON_MIN <= lon_f <= LON_MAX):
                out_of_bounds.append((r["ins_id"], lat_f, lon_f))
            lat, lon = f"{lat_f:.6f}", f"{lon_f:.6f}"
        # The legal entity behind the institution: its own JAR code, else (a branch
        # with an internal ŠVIS code) the parent institution's JAR code.
        legal = r["ins_kodas"] if is_jar(r["ins_kodas"]) else ""
        parent = ""
        if school.get("ar_padalinys") == "1" and school.get("ins_motin_id") not in ("", r["ins_id"]):
            parent = school["ins_motin_id"]
            if not legal and is_jar(school.get("ins_motin_kodas", "")):
                legal = school["ins_motin_kodas"]
        sen = r["sen_kodas"] if re.fullmatch(r"[0-9]{4}", r["sen_kodas"] or "") else ""
        kind = r["ins_pagr_tipas"] or r["ins_pask_grupe"]
        out.append({
            "ins_id": r["ins_id"],
            "ins_kodas": r["ins_kodas"],
            "name": r["ins_pavad_lt"].strip(),
            "registered": iso_date(r["ins_ireg_smir_data"]),
            "deregistered": iso_date(r["ins_isreg_smir_data"]),
            "type": kind,
            "type_slug": slugify(kind) if kind else "",
            "ownership": school.get("prk_pavad", ""),
            "sav_kodas": r["savivaldybe"],
            "sen_kodas": sen,
            "gyv_kodas": r["gyv_kodas"],
            "street": school.get("gat_pavadinimas", ""),
            "house": school.get("adr_namas", ""),
            "lat": lat,
            "long": lon,
            "legal_entity": legal,
            "parent": parent,
        })
    if out_of_bounds:
        print(f"ERROR: {len(out_of_bounds)} point(s) outside Lithuania's bbox (axis swap?): "
              f"{out_of_bounds[:5]}", file=sys.stderr)
        sys.exit(1)
    write_csv(args.output, list(out[0].keys()), out)
    types = Counter((r["type_slug"], r["type"]) for r in out if r["type_slug"])
    write_csv(args.types, ["slug", "label"],
              [{"slug": s, "label": l} for (s, l) in sorted(types)])


# --- health --------------------------------------------------------------------

# "Klaipėdos apskr. Klaipėdos miesto sav. Klaipėdos m. Taikos pr. 5-207"
# "Marijampolės apskr., Vilkaviškio r. sav., Pilviškių sen., Ramoniškių k., Rausvės g. 4"
APSKR_RE = re.compile(r"^.*?\bapskr\.,?\s*")
SAV_RE = re.compile(r"^([^,]+?sav\.)")
SEN_RE = re.compile(r"^[^,]+?sav\.,?\s+([^,]+?\s+sen\.)")


def health(args) -> None:
    names = AdminNames(args.municipalities, args.elderships)
    institutions = read_csv(args.institutions)
    known = {r["spi_id"] for r in institutions}
    out = []
    for r in institutions:
        out.append({
            "spi_id": r["spi_id"],
            "ja_kodas": r["ja_kodas"] if is_jar(r["ja_kodas"]) else "",
            "name": r["ja_pavadinimas"].strip(),
            "legal_form": r["teisine_forma"],
            "licence_nr": r["licencijos_nr"],
            "licence_date": iso_date(r["licencijos_isdavimo_data"]),
        })
    write_csv(args.output, list(out[0].keys()), out)

    sites, orphans, unparsed = [], 0, []
    for r in read_csv(args.addresses):
        if r["istaiga"] not in known:
            orphans += 1  # address of an institution no longer in the licence list
            continue
        text = re.sub(r"\s+", " ", r["adresas"]).strip()
        sav = sen = ""
        rest = APSKR_RE.sub("", text)
        m = SAV_RE.match(rest)
        if m:
            sav = names.municipality(m.group(1))
            s = SEN_RE.match(rest)
            if s and sav:
                sen = names.eldership(sav, s.group(1))
        else:
            names.misses["unparsed"] += 1
            unparsed.append(text)
        sites.append({
            "site_id": r["spi_veiklos_adreso_id"].lower(),
            "spi_id": r["istaiga"],
            "address": text,
            "sav_kodas": sav,
            "sen_kodas": sen,
        })
    print(f"health: {orphans} activity addresses of unlisted institutions skipped", file=sys.stderr)
    names.report("health")
    if unparsed:
        print(f"health: unparsed address samples {unparsed[:5]}", file=sys.stderr)
    write_csv(args.sites, list(sites[0].keys()), sites)


# --- heritage ------------------------------------------------------------------

# the KVR "valuable property" flags (vertingųjų savybių pobūdis), column -> label
VALUE_FLAGS = ["archeologinis", "architektūrinis", "dailės", "etnokultūrinis",
               "inžinerinis", "istorinis", "kraštovaizdžio", "kultūrinės_raiškos",
               "memorialinis", "povandeninis", "mitologinis", "sakralinis",
               "urbanistinis", "želdynų"]


def heritage(args) -> None:
    from openpyxl import load_workbook

    names = AdminNames(args.municipalities, args.elderships)
    wb = load_workbook(args.input, read_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    out, values, codes = [], [], set()
    concepts = {"statuses": Counter(), "significance": Counter(), "kinds": Counter()}

    def clean(v) -> str:
        return "" if v in (None, "NULL") else str(v).strip()

    for row in it:
        r = dict(zip(header, row))
        code = clean(r["unikalus_kodas"])
        if not code:
            continue
        codes.add(code)
        status, signif, kind = clean(r["statusas"]), clean(r["reikšmingumas"]), clean(r["potipis"])
        for scheme, label in (("statuses", status), ("significance", signif), ("kinds", kind)):
            if label:
                concepts[scheme][label] += 1
        sav = names.municipality(clean(r["savivaldybė"]))
        sen = names.eldership(sav, clean(r["seniūnija"]))
        parent = clean(r["priklauso_kompleksui"])
        out.append({
            "code": code,
            "parent": parent if parent and parent != code else "",
            "name": clean(r["pavadinimas"]),
            "registered": iso_date(r["registracijos_data"]),
            "status_slug": slugify(status) if status else "",
            "significance_slug": slugify(signif) if signif else "",
            "kind_slug": slugify(kind) if kind else "",
            "classification": clean(r["klasifikacija"]),
            "sav_kodas": sav,
            "sen_kodas": sen,
        })
        for flag in VALUE_FLAGS:
            if str(r.get(flag)) == "1":
                values.append({"code": code, "value_slug": slugify(flag)})
    # a complex member whose complex is absent from the dump would dangle
    dangling = 0
    for r in out:
        if r["parent"] and r["parent"] not in codes:
            r["parent"] = ""
            dangling += 1
    print(f"heritage: {dangling} complex links to objects absent from the dump dropped",
          file=sys.stderr)
    names.report("heritage")
    write_csv(args.output, list(out[0].keys()), out)
    write_csv(args.values, ["code", "value_slug"], values)
    taxonomy = []
    for scheme, counter in concepts.items():
        for label in sorted(counter):
            taxonomy.append({"scheme": f"heritage-{scheme}", "slug": slugify(label), "label": label})
    for flag in VALUE_FLAGS:
        label = flag.replace("_", " ").capitalize()
        taxonomy.append({"scheme": "heritage-values", "slug": slugify(flag), "label": label})
    write_csv(args.taxonomy, ["scheme", "slug", "label"], taxonomy)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("education")
    p.add_argument("--institutions", required=True)
    p.add_argument("--schools", required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--types", required=True)
    p.set_defaults(func=education)

    p = sub.add_parser("health")
    p.add_argument("--institutions", required=True)
    p.add_argument("--addresses", required=True)
    p.add_argument("--municipalities", required=True)
    p.add_argument("--elderships", required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--sites", required=True)
    p.set_defaults(func=health)

    p = sub.add_parser("heritage")
    p.add_argument("--input", required=True)
    p.add_argument("--municipalities", required=True)
    p.add_argument("--elderships", required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--values", required=True)
    p.add_argument("--taxonomy", required=True)
    p.set_defaults(func=heritage)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
